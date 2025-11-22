import argparse
import csv
import sys
import re
from pathlib import Path
from collections import Counter

import numpy as np

try:
    import mne
    mne.set_log_level("WARNING")
except Exception:
    mne = None

try:
    from scipy.io import loadmat
except Exception:
    loadmat = None

try:
    import matplotlib.pyplot as plt
except Exception:
    plt = None

try:
    import seaborn as sns
    if sns is not None:
        sns.set_theme(style="whitegrid")
except Exception:
    sns = None

try:
    from openpyxl import Workbook
except Exception:
    Workbook = None

DEFAULT_DATA_DIR = "data"
DEFAULT_OUTPUT_DIR = "."

def format_seconds(total_seconds: float) -> str:
    if total_seconds is None:
        return "unknown"
    total_seconds = float(total_seconds)
    h = int(total_seconds // 3600)
    m = int((total_seconds % 3600) // 60)
    s = total_seconds - h * 3600 - m * 60
    return f"{h:02d}:{m:02d}:{s:06.3f}"

def as_float(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        if isinstance(x, np.generic):
            return float(np.asarray(x).item())
    except Exception:
        pass
    if isinstance(x, (list, tuple)):
        return as_float(x[0]) if x else None
    if isinstance(x, np.ndarray):
        if x.size == 1:
            try:
                return float(x.ravel()[0])
            except Exception:
                return None
    try:
        return float(x)
    except Exception:
        return None

def get_field(obj, key):
    if obj is None:
        return None
    if isinstance(obj, dict):
        return obj.get(key)
    try:
        if hasattr(obj, key):
            return getattr(obj, key)
    except Exception:
        pass
    if isinstance(obj, np.void):
        names = obj.dtype.names or []
        if key in names:
            val = obj[key]
            if isinstance(val, np.ndarray) and val.size == 1:
                try:
                    return val.item()
                except Exception:
                    return val
            return val
    return None

def to_ndarray(x):
    if isinstance(x, np.ndarray):
        if x.dtype == object and x.size == 1:
            try:
                return np.asarray(x.item())
            except Exception:
                return x
        return x
    try:
        return np.asarray(x)
    except Exception:
        return None

def to_str(x):
    if x is None:
        return None
    if isinstance(x, str):
        return x
    if isinstance(x, bytes):
        try:
            return x.decode("utf-8", errors="ignore")
        except Exception:
            return str(x)
    if isinstance(x, np.ndarray):
        try:
            if x.dtype.kind in ("U", "S"):
                return "".join(x.astype(str).tolist())
            if x.size == 1:
                return to_str(x.item())
            return ",".join([to_str(el) or "" for el in x.ravel()])
        except Exception:
            return str(x)
    try:
        return str(x)
    except Exception:
        return None

def analyze_raw(raw):
    sfreq = as_float(raw.info.get("sfreq"))
    n_times = int(raw.n_times)
    duration_sec = (n_times / sfreq) if sfreq and sfreq > 0 else None
    n_epochs = 0
    counts = Counter()
    try:
        ann = getattr(raw, "annotations", None)
        if ann is not None and len(ann) > 0:
            for d in ann.description:
                if d is not None and str(d) != "":
                    counts[str(d)] += 1
        else:
            try:
                events = mne.find_events(raw, stim_channel="auto", shortest_event=1, verbose=False)
                if events is not None and len(events) > 0:
                    code_counts = Counter(events[:, 2].tolist())
                    for code, cnt in code_counts.items():
                        counts[str(code)] += cnt
            except Exception:
                pass
    except Exception:
        pass
    has_labels = sum(counts.values()) > 0
    n_channels = len(raw.ch_names)
    return {
        "type": "raw",
        "sfreq": sfreq,
        "n_channels": n_channels,
        "duration_sec": duration_sec,
        "n_epochs": n_epochs,
        "labels_counts": dict(counts),
        "has_labels": has_labels,
    }

def analyze_epochs(epochs):
    sfreq = as_float(epochs.info.get("sfreq"))
    n_epochs = len(epochs)
    n_times = int(epochs.n_times)
    epoch_dur = (n_times / sfreq) if sfreq and sfreq > 0 else None
    duration_sec = (n_epochs * epoch_dur) if epoch_dur is not None else None
    counts = Counter()
    try:
        event_id = getattr(epochs, "event_id", {}) or {}
        inv_map = {v: k for k, v in event_id.items()} if isinstance(event_id, dict) else {}
        events = getattr(epochs, "events", None)
        if events is not None:
            for code in events[:, 2].tolist():
                label = inv_map.get(int(code), str(int(code)))
                counts[label] += 1
    except Exception:
        pass
    has_labels = sum(counts.values()) > 0
    n_channels = len(epochs.ch_names)
    return {
        "type": "epochs",
        "sfreq": sfreq,
        "n_channels": n_channels,
        "duration_sec": duration_sec,
        "n_epochs": n_epochs,
        "labels_counts": dict(counts),
        "has_labels": has_labels,
    }

def ensure_output_dir(path: Path):
    try:
        path.mkdir(parents=True, exist_ok=True)
    except Exception:
        pass

def save_csv_summaries(results, total_counts: Counter, output_dir: Path, prefix: str):
    files_csv_path = output_dir / f"{prefix}_files.csv"
    labels_csv_path = output_dir / f"{prefix}_labels.csv"
    fields = [
        "file",
        "format",
        "type",
        "sfreq",
        "n_channels",
        "duration_sec",
        "duration_str",
        "n_epochs",
        "has_labels",
        "labels_total",
    ]
    try:
        with files_csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(fields)
            for r in results:
                dur = r.get("duration_sec")
                labels_total = sum((Counter(r.get("labels_counts") or {})).values())
                writer.writerow([
                    r.get("file"),
                    r.get("format"),
                    r.get("type"),
                    r.get("sfreq"),
                    r.get("n_channels"),
                    ("" if dur is None else f"{float(dur):.6f}"),
                    format_seconds(dur),
                    r.get("n_epochs"),
                    r.get("has_labels"),
                    labels_total,
                ])
    except Exception:
        pass
    try:
        with labels_csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["label", "count"])
            for lab, cnt in sorted(total_counts.items(), key=lambda x: (-x[1], x[0])):
                writer.writerow([lab, cnt])
    except Exception:
        pass
    return str(files_csv_path), str(labels_csv_path)

def plot_summaries(results, total_counts: Counter, output_dir: Path, prefix: str):
    saved = []
    if plt is None:
        return saved
    if total_counts and sum(total_counts.values()) > 0:
        top_items = sorted(total_counts.items(), key=lambda x: (-x[1], x[0]))[:30]
        labels, counts = zip(*top_items)
        fig, ax = plt.subplots(figsize=(10, max(4, len(labels) * 0.35)))
        try:
            if sns is not None:
                sns.barplot(x=list(counts), y=list(labels), orient="h", ax=ax, color="#4C78A8")
            else:
                ax.barh(labels, counts, color="#4C78A8")
        except Exception:
            ax.barh(labels, counts, color="#4C78A8")
        ax.set_title("Распределение меток (топ-30)")
        ax.set_xlabel("Количество")
        ax.set_ylabel("Метка")
        plt.tight_layout()
        out_path = output_dir / f"{prefix}_labels.png"
        try:
            fig.savefig(out_path, dpi=150)
            saved.append(str(out_path))
        except Exception:
            pass
        plt.close(fig)
    file_durs = []
    for r in results:
        d = r.get("duration_sec")
        if d is not None:
            try:
                file_durs.append((r.get("file"), float(d)))
            except Exception:
                continue
    if file_durs:
        file_durs.sort(key=lambda x: x[1], reverse=True)
        files, durations = zip(*file_durs)
        fig, ax = plt.subplots(figsize=(10, max(4, len(files) * 0.35)))
        try:
            if sns is not None:
                sns.barplot(x=list(durations), y=list(files), orient="h", ax=ax, color="#72B7B2")
            else:
                ax.barh(files, durations, color="#72B7B2")
        except Exception:
            ax.barh(files, durations, color="#72B7B2")
        ax.set_title("Длительность записей по файлам (сек)")
        ax.set_xlabel("Секунды")
        ax.set_ylabel("Файл")
        plt.tight_layout()
        out_path = output_dir / f"{prefix}_durations.png"
        try:
            fig.savefig(out_path, dpi=150)
            saved.append(str(out_path))
        except Exception:
            pass
        plt.close(fig)
    file_epochs = []
    for r in results:
        try:
            file_epochs.append((r.get("file"), int(r.get("n_epochs") or 0)))
        except Exception:
            file_epochs.append((r.get("file"), 0))
    if file_epochs:
        file_epochs.sort(key=lambda x: x[1], reverse=True)
        files, epochs_vals = zip(*file_epochs)
        fig, ax = plt.subplots(figsize=(10, max(4, len(files) * 0.35)))
        try:
            if sns is not None:
                sns.barplot(x=list(epochs_vals), y=list(files), orient="h", ax=ax, color="#54A24B")
            else:
                ax.barh(files, epochs_vals, color="#54A24B")
        except Exception:
            ax.barh(files, epochs_vals, color="#54A24B")
        ax.set_title("Количество эпох по файлам")
        ax.set_xlabel("Эпохи")
        ax.set_ylabel("Файл")
        plt.tight_layout()
        out_path = output_dir / f"{prefix}_epochs.png"
        try:
            fig.savefig(out_path, dpi=150)
            saved.append(str(out_path))
        except Exception:
            pass
        plt.close(fig)
    return saved

def export_epochs_to_excel(epochs_list, output_dir: Path, prefix: str):
    if Workbook is None or not epochs_list:
        return None
    def _safe_sheet_name(name: str) -> str:
        bad = set('[]:*?/\\')
        s = ''.join('_' if ch in bad else ch for ch in (name or 'Sheet'))
        return (s or 'Sheet')[:31]
    wb = Workbook()
    try:
        wb.remove(wb.active)
    except Exception:
        pass
    used = set()
    for file_name, epochs in epochs_list:
        try:
            sfreq = as_float(epochs.info.get("sfreq")) or 0.0
            n_samp = int(epochs.n_times)
            sec_per_epoch = (n_samp / sfreq) if sfreq and sfreq > 0 else None
            event_id = getattr(epochs, "event_id", {}) or {}
            inv_map = {v: k for k, v in event_id.items()} if isinstance(event_id, dict) else {}
            events = getattr(epochs, "events", None)
            sheet = _safe_sheet_name(file_name)
            base = sheet
            idx = 2
            while sheet in used:
                sheet = _safe_sheet_name(f"{base}_{idx}")
                idx += 1
            used.add(sheet)
            ws = wb.create_sheet(title=sheet)
            ws.append(["file", "epoch_index", "samples_per_epoch", "seconds_per_epoch", "label_code", "label_name"])
            n_ep = len(epochs)
            for i in range(n_ep):
                code = None
                name = None
                try:
                    if events is not None:
                        code = int(events[i, 2])
                        name = inv_map.get(code, str(code))
                except Exception:
                    pass
                ws.append([
                    file_name,
                    i + 1,
                    n_samp,
                    (None if sec_per_epoch is None else float(sec_per_epoch)),
                    code,
                    name,
                ])
        except Exception:
            continue
    out_path = output_dir / f"{prefix}_epochs.xlsx"
    try:
        wb.save(out_path)
        return str(out_path)
    except Exception:
        return None

def loadmat_safely(f: str):
    if loadmat is None:
        return None
    try:
        return loadmat(f, simplify_cells=True)
    except TypeError:
        try:
            return loadmat(f, struct_as_record=False, squeeze_me=True)
        except Exception:
            return None
    except Exception:
        return None

def analyze_mat(path: Path):
    info = {
        "type": "mat",
        "sfreq": None,
        "n_channels": None,
        "duration_sec": None,
        "n_epochs": 0,
        "labels_counts": {},
        "has_labels": False,
    }
    mm = loadmat_safely(str(path))
    if mm is None:
        return info
    sfreq = None
    n_channels = None
    duration_sec = None
    n_epochs = 0
    counts = Counter()
    samples_per_epoch = None
    labels_per_epoch = None
    eeg = mm.get("EEG") if "EEG" in mm else None
    data = None
    if eeg is not None:
        data = get_field(eeg, "data")
        if data is None:
            data = get_field(eeg, "x")
        sr = get_field(eeg, "srate")
        if sr is None:
            sr = get_field(eeg, "fs")
        if sr is None:
            sr = get_field(eeg, "Fs")
        sfreq = as_float(sr)
    if data is None:
        if "data" in mm:
            data = mm.get("data")
        elif "X" in mm:
            data = mm.get("X")
        elif "signals" in mm:
            data = mm.get("signals")
        elif "x" in mm:
            data = mm.get("x")
    if sfreq is None:
        for k in ("fs", "Fs", "srate", "sampling_rate", "sr"):
            if k in mm:
                sfreq = as_float(mm.get(k))
                if sfreq:
                    break
    arr = to_ndarray(data) if data is not None else None
    if isinstance(arr, np.ndarray) and arr.size > 0 and arr.ndim in (2, 3):
        if arr.ndim == 2:
            n_channels, n_times = arr.shape
            n_epochs = 0
            if sfreq and sfreq > 0:
                duration_sec = n_times / sfreq
            try:
                samples_per_epoch = int(n_times)
            except Exception:
                samples_per_epoch = None
        else:
            n_channels, n_times, n_epochs = arr.shape
            if sfreq and sfreq > 0:
                duration_sec = (n_times * n_epochs) / sfreq
            try:
                samples_per_epoch = int(n_times)
            except Exception:
                samples_per_epoch = None
        n_channels = int(n_channels)
    # Сопоставление меток ТОЛЬКО через EEG.event(i).epoch (1-based -> 0-based)
    if eeg is not None and n_epochs and isinstance(n_epochs, (int, np.integer)) and int(n_epochs) > 0:
        events = get_field(eeg, "event")
        if events is not None:
            try:
                tmp = [[] for _ in range(int(n_epochs))]
                if isinstance(events, list):
                    ev_iter = events
                elif isinstance(events, np.ndarray):
                    ev_iter = events.ravel().tolist()
                else:
                    ev_iter = []
                for ev in ev_iter:
                    ep = get_field(ev, "epoch") if isinstance(ev, (dict, np.void)) or hasattr(ev, "__dict__") else None
                    if ep is None:
                        continue
                    try:
                        idx = int(as_float(ep)) - 1
                    except Exception:
                        idx = None
                    if idx is None or idx < 0 or idx >= int(n_epochs):
                        continue
                    t = get_field(ev, "type") if isinstance(ev, (dict, np.void)) or hasattr(ev, "__dict__") else ev
                    lab = to_str(t)
                    if lab:
                        tmp[idx].append(lab)
                if any(len(x) > 0 for x in tmp):
                    labels_per_epoch = tmp
            except Exception:
                pass
    # Подсчёт распределения меток: по labels_per_epoch, если удалось собрать, иначе по всем событиям
    if labels_per_epoch:
        flat = []
        for lab in labels_per_epoch:
            if not lab:
                continue
            if isinstance(lab, (list, tuple, np.ndarray)):
                try:
                    seq = np.asarray(lab).ravel().tolist()
                except Exception:
                    seq = lab
                for el in seq:
                    s = to_str(el)
                    if s:
                        flat.append(s)
            else:
                s = to_str(lab)
                if s:
                    flat.append(s)
        if flat:
            counts.update([str(l) for l in flat if str(l) not in ("", "nan", "None")])
    else:
        if eeg is not None:
            events = get_field(eeg, "event")
            if events is not None:
                try:
                    if isinstance(events, list):
                        ev_iter = events
                    elif isinstance(events, np.ndarray):
                        ev_iter = events.ravel().tolist()
                    else:
                        ev_iter = []
                    for ev in ev_iter:
                        t = get_field(ev, "type") if isinstance(ev, (dict, np.void)) or hasattr(ev, "__dict__") else ev
                        lab = to_str(t)
                        if lab:
                            counts[lab] += 1
                except Exception:
                    pass
    info.update({
        "sfreq": sfreq,
        "n_channels": n_channels,
        "duration_sec": duration_sec,
        "n_epochs": int(n_epochs) if isinstance(n_epochs, (int, np.integer)) else 0,
        "labels_counts": dict(counts),
        "has_labels": sum(counts.values()) > 0,
        "samples_per_epoch": (int(samples_per_epoch) if isinstance(samples_per_epoch, (int, np.integer)) or (samples_per_epoch is not None and str(type(samples_per_epoch)).endswith("int'>")) else None),
        "labels_per_epoch": labels_per_epoch,
    })
    return info

def export_mat_epochs_to_excel(results, output_dir: Path, prefix: str):
    if Workbook is None:
        return None
    def _safe_sheet_name(name: str) -> str:
        bad = set('[]:*?/\\')
        s = ''.join('_' if ch in bad else ch for ch in (name or 'Sheet'))
        return (s or 'Sheet')[:31]
    wb = Workbook()
    try:
        wb.remove(wb.active)
    except Exception:
        pass
    used = set()
    any_written = False
    for r in results:
        try:
            if r.get("type") != "mat":
                continue
            n_epochs = int(r.get("n_epochs") or 0)
            sfreq = as_float(r.get("sfreq")) or None
            n_samp = r.get("samples_per_epoch")
            if not n_epochs or n_samp is None:
                continue
            try:
                n_samp = int(n_samp)
            except Exception:
                continue
            sec_per_epoch = (float(n_samp) / float(sfreq)) if (sfreq and sfreq > 0) else None
            sheet = _safe_sheet_name(r.get("file"))
            base = sheet
            idx = 2
            while sheet in used:
                sheet = _safe_sheet_name(f"{base}_{idx}")
                idx += 1
            used.add(sheet)
            ws = wb.create_sheet(title=sheet)
            ws.append(["file", "epoch_index", "samples_per_epoch", "seconds_per_epoch", "label_code", "label_name"])
            labels = r.get("labels_per_epoch") or []
            for i in range(n_epochs):
                name = None
                code = None
                if i < len(labels):
                    raw_lab = labels[i]
                    if raw_lab is None:
                        names = []
                    elif isinstance(raw_lab, (list, tuple, np.ndarray)):
                        try:
                            seq = np.asarray(raw_lab).ravel().tolist()
                        except Exception:
                            seq = raw_lab
                        names = [to_str(el) for el in seq if to_str(el)]
                    else:
                        n = to_str(raw_lab)
                        names = [n] if n else []
                    if names:
                        name = "; ".join(names)
                        codes = []
                        for nm in names:
                            try:
                                m = re.search(r"(\d+)", nm)
                                if m:
                                    codes.append(str(int(m.group(1))))
                            except Exception:
                                continue
                        code = "; ".join(codes) if codes else None
                ws.append([
                    r.get("file"),
                    i + 1,
                    n_samp,
                    (None if sec_per_epoch is None else float(sec_per_epoch)),
                    code,
                    name,
                ])
            any_written = True
        except Exception:
            continue
    if not any_written:
        return None
    out_path = output_dir / f"{prefix}_epochs.xlsx"
    try:
        wb.save(out_path)
        return str(out_path)
    except Exception:
        return None

def read_with_mne(path: Path):
    if mne is None:
        return None, None
    f = str(path)
    lower = f.lower()
    try:
        if lower.endswith(".edf"):
            raw = mne.io.read_raw_edf(f, preload=False, verbose=False)
            return "raw", raw
        if lower.endswith(".bdf"):
            raw = mne.io.read_raw_bdf(f, preload=False, verbose=False)
            return "raw", raw
        if lower.endswith(".vhdr"):
            raw = mne.io.read_raw_brainvision(f, preload=False, verbose=False)
            return "raw", raw
        if lower.endswith(".cnt"):
            raw = mne.io.read_raw_cnt(f, preload=False, verbose=False)
            return "raw", raw
        if lower.endswith(".eeg"):
            vhdr = path.with_suffix(".vhdr")
            if vhdr.exists():
                raw = mne.io.read_raw_brainvision(str(vhdr), preload=False, verbose=False)
                return "raw", raw
        if lower.endswith(".fif"):
            try:
                raw = mne.io.read_raw_fif(f, preload=False, verbose=False)
                return "raw", raw
            except Exception:
                try:
                    epochs = mne.read_epochs(f, preload=False, verbose=False)
                    return "epochs", epochs
                except Exception:
                    return None, None
        if lower.endswith(".set"):
            try:
                epochs = mne.read_epochs_eeglab(f, verbose=False)
                return "epochs", epochs
            except Exception:
                raw = mne.io.read_raw_eeglab(f, preload=False, verbose=False)
                return "raw", raw
    except Exception:
        return None, None
    return None, None

def analyze_file(path: Path, return_obj: bool = False):
    ext = path.suffix.lower()
    if ext in {".edf", ".bdf", ".vhdr", ".cnt", ".eeg", ".fif", ".set"}:
        kind, obj = read_with_mne(path)
        if kind == "raw" and obj is not None:
            res = analyze_raw(obj)
            res.update({"format": ext, "file": path.name, "path": str(path)})
            return (res, obj) if return_obj else res
        if kind == "epochs" and obj is not None:
            res = analyze_epochs(obj)
            res.update({"format": ext, "file": path.name, "path": str(path)})
            return (res, obj) if return_obj else res
        unk = {"type": "unknown", "format": ext, "file": path.name, "path": str(path), "sfreq": None, "n_channels": None, "duration_sec": None, "n_epochs": 0, "labels_counts": {}, "has_labels": False}
        return (unk, None) if return_obj else unk
    if ext == ".mat":
        res = analyze_mat(path)
        res.update({"format": ext, "file": path.name, "path": str(path)})
        return (res, None) if return_obj else res
    skipped = {"type": "skipped", "format": ext, "file": path.name, "path": str(path), "sfreq": None, "n_channels": None, "duration_sec": None, "n_epochs": 0, "labels_counts": {}, "has_labels": False}
    return (skipped, None) if return_obj else skipped

def gather_files(data_dir: Path, recursive: bool = False):
    exts = {".edf", ".bdf", ".fif", ".vhdr", ".set", ".cnt", ".eeg", ".mat"}
    files = []
    if recursive:
        for p in sorted(data_dir.rglob("*")):
            if p.is_file() and p.suffix.lower() in exts:
                files.append(p)
    else:
        for p in sorted(data_dir.iterdir()):
            if p.is_file() and p.suffix.lower() in exts:
                files.append(p)
    return files

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("pos_data_dir", nargs="?", default=None, help="Папка с данными (по умолчанию 'data')")
    parser.add_argument("pos_output_dir", nargs="?", default=None, help="Папка вывода (по умолчанию '.')")
    parser.add_argument("--data-dir", type=str, default=None, help="Папка с данными")
    parser.add_argument("--recursive", action="store_true", help="Рекурсивный обход подкаталогов")
    parser.add_argument("--output-dir", type=str, default=None, help="Папка для сохранения CSV и графиков")
    parser.add_argument("--output-csv-prefix", type=str, default="eeg_analysis", help="Префикс имен CSV и графиков")
    parser.add_argument("--no-plots", action="store_true", help="Отключить построение графиков")
    args = parser.parse_args()

    data_dir_str = args.data_dir or args.pos_data_dir or DEFAULT_DATA_DIR
    out_dir_str = args.output_dir or args.pos_output_dir or DEFAULT_OUTPUT_DIR

    data_dir = Path(data_dir_str)
    out_dir = Path(out_dir_str)
    ensure_output_dir(out_dir)
    if not data_dir.exists() or not data_dir.is_dir():
        print(f"Папка не найдена: {data_dir}")
        sys.exit(1)

    files = gather_files(data_dir, recursive=args.recursive)
    if not files:
        print(f"EEG файлы не найдены в папке {data_dir}")
        sys.exit(0)

    results = []
    total_duration = 0.0
    total_epochs = 0
    total_counts = Counter()
    epochs_exports = []

    for fp in files:
        try:
            out = analyze_file(fp, return_obj=True)
            if isinstance(out, tuple) and len(out) == 2:
                res, obj = out
            else:
                res, obj = out, None
        except Exception as e:
            res = {"type": "error", "format": fp.suffix.lower(), "file": fp.name, "path": str(fp), "sfreq": None, "n_channels": None, "duration_sec": None, "n_epochs": 0, "labels_counts": {}, "has_labels": False, "error": str(e)}
            obj = None
        results.append(res)
        if obj is not None and res.get("type") == "epochs":
            epochs_exports.append((fp.name, obj))

    print("ФАЙЛЫ:")
    for r in results:
        dur = r.get("duration_sec")
        print(f"- {r.get('file')} | формат: {r.get('format')} | тип: {r.get('type')} | fs: {r.get('sfreq')} | каналы: {r.get('n_channels')} | длительность: {format_seconds(dur)} | эпох: {r.get('n_epochs')} | метки: {sum((Counter(r.get('labels_counts') or {})).values())}")
        if dur is not None:
            total_duration += float(dur)
        total_epochs += int(r.get("n_epochs") or 0)
        total_counts.update(r.get("labels_counts") or {})

    print("")
    print("ИТОГО:")
    print(f"- Совокупная длительность: {format_seconds(total_duration)} ({total_duration:.3f} с)")
    print(f"- Общее число эпох: {total_epochs}")
    if total_counts:
        print("- Распределение меток:")
        for k, v in sorted(total_counts.items(), key=lambda x: (-x[1], x[0])):
            print(f"  {k}: {v}")
    else:
        print("- Метки не найдены")

    files_csv, labels_csv = save_csv_summaries(results, total_counts, out_dir, args.output_csv_prefix)
    print("")
    print(f"CSV файлы сохранены:\n- {files_csv}\n- {labels_csv}")

    if not args.no_plots:
        saved_plots = plot_summaries(results, total_counts, out_dir, args.output_csv_prefix)
        if saved_plots:
            print("Графики сохранены:")
            for pth in saved_plots:
                print(f"- {pth}")

    if epochs_exports:
        if Workbook is None:
            print("Пропущен экспорт Excel по эпохам: пакет openpyxl не установлен")
        else:
            xlsx_path = export_epochs_to_excel(epochs_exports, out_dir, args.output_csv_prefix)
            if xlsx_path:
                print(f"Excel по эпохам сохранен: {xlsx_path}")
    else:
        if Workbook is None:
            print("Пропущен экспорт Excel по эпохам: пакет openpyxl не установлен")
        else:
            xlsx_path = export_mat_epochs_to_excel(results, out_dir, args.output_csv_prefix)
            if xlsx_path:
                print(f"Excel по эпохам сохранен: {xlsx_path}")

if __name__ == "__main__":
    main()
