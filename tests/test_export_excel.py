import os
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

import analyze_eeg_mne as mod


class DummyEpochs:
    def __init__(self, n_epochs: int, n_times: int, sfreq: float, events=None, event_id=None, ch_names=None):
        self._n_epochs = int(n_epochs)
        self.n_times = int(n_times)
        self.info = {"sfreq": float(sfreq)}
        self.events = events
        self.event_id = event_id or {}
        self.ch_names = ch_names or ["Cz"]

    def __len__(self):
        return self._n_epochs


def test_export_epochs_to_excel_with_dummy_epochs(tmp_path: Path):
    n_epochs = 3
    n_times = 200
    sfreq = 100.0
    events = np.array([
        [0, 0, 1],
        [0, 0, 2],
        [0, 0, 1],
    ], dtype=int)
    event_id = {"A": 1, "B": 2}
    dummy = DummyEpochs(n_epochs=n_epochs, n_times=n_times, sfreq=sfreq, events=events, event_id=event_id)

    out = mod.export_epochs_to_excel([("dummy.set", dummy)], tmp_path, "eeg_analysis")
    assert out is not None
    assert os.path.exists(out)

    wb = load_workbook(out)
    assert "dummy.set" in wb.sheetnames
    ws = wb["dummy.set"]
    # header + n_epochs rows
    assert ws.max_row == 1 + n_epochs
    # Check header
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert hdr == ["file", "epoch_index", "samples_per_epoch", "seconds_per_epoch", "label_code", "label_name"]
    # Check one row
    row2 = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row2[0] == "dummy.set"
    assert row2[1] == 1
    assert row2[2] == n_times
    assert row2[3] == n_times / sfreq


def test_export_mat_epochs_to_excel_from_results(tmp_path: Path):
    results = [
        {
            "type": "mat",
            "file": "example.mat",
            "n_epochs": 4,
            "sfreq": 200.0,
            "samples_per_epoch": 250,
        },
        {
            "type": "raw",
            "file": "skip.edf",
        },
    ]
    out = mod.export_mat_epochs_to_excel(results, tmp_path, "eeg_analysis")
    assert out is not None
    assert os.path.exists(out)

    wb = load_workbook(out)
    assert "example.mat" in wb.sheetnames
    ws = wb["example.mat"]
    # header + n_epochs rows
    assert ws.max_row == 1 + 4
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert hdr == ["file", "epoch_index", "samples_per_epoch", "seconds_per_epoch", "label_code", "label_name"]
    row_last = [c.value for c in next(ws.iter_rows(min_row=5, max_row=5))]
    assert row_last[1] == 4
    assert row_last[2] == 250
    assert row_last[3] == 250 / 200.0
